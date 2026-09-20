from __future__ import annotations

import csv
import hashlib
import io
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook


SAMPLE_ROWS = 5


@dataclass
class ParsedSpreadsheet:
    source_file_name: str
    source_file_hash: str
    headers: list[str]
    rows: list[dict[str, Optional[str]]]
    samples: dict[str, list[str]]
    stats: dict[str, int] = field(default_factory=dict)


def _cell_to_str(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text if text else None


def _normalize_headers(raw: list[Any]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(raw):
        base = _cell_to_str(value) or f"Column {index + 1}"
        count = seen.get(base, 0)
        seen[base] = count + 1
        headers.append(base if count == 0 else f"{base} ({count + 1})")
    return headers


def _estimate_stats(
    headers: list[str], rows: list[dict[str, Optional[str]]]
) -> dict[str, int]:
    section_key = next((h for h in headers if h.lower().startswith("section name")), None)
    item_key = next((h for h in headers if h.lower().startswith("item name")), None)

    sections: set[str] = set()
    items: set[tuple[str, str]] = set()
    for row in rows:
        section = (row.get(section_key) or "").strip() if section_key else ""
        item = (row.get(item_key) or "").strip() if item_key else ""
        if section:
            sections.add(section)
        if section or item:
            items.add((section, item or ""))

    return {
        "totalRows": len(rows),
        "estSections": len(sections) if sections else (1 if rows else 0),
        "estLineItems": len(items) if items else len(rows),
        "estComments": len(rows),
    }


def _build_samples(
    headers: list[str], rows: list[dict[str, Optional[str]]]
) -> dict[str, list[str]]:
    samples: dict[str, list[str]] = {header: [] for header in headers}
    for row in rows[:SAMPLE_ROWS]:
        for header in headers:
            value = row.get(header)
            samples[header].append(value if value is not None else "")
    return samples


def _parse_csv(content: bytes) -> tuple[list[str], list[dict[str, Optional[str]]]]:
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    try:
        header_row = next(reader)
    except StopIteration as exc:
        raise ValueError("CSV file is empty") from exc

    headers = _normalize_headers(header_row)
    rows: list[dict[str, Optional[str]]] = []
    for raw in reader:
        if not any(_cell_to_str(cell) for cell in raw):
            continue
        padded = list(raw) + [None] * max(0, len(headers) - len(raw))
        rows.append(
            {headers[i]: _cell_to_str(padded[i]) for i in range(len(headers))}
        )
    return headers, rows


def _parse_excel(content: bytes) -> tuple[list[str], list[dict[str, Optional[str]]]]:
    # Spectora often ships OOXML workbooks with a .xls extension.
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp.write(content)
        tmp_path = Path(tmp.name)

    try:
        workbook = load_workbook(tmp_path, read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        try:
            header_row = next(iterator)
        except StopIteration as exc:
            raise ValueError("Excel file is empty") from exc

        headers = _normalize_headers(list(header_row))
        rows: list[dict[str, Optional[str]]] = []
        for raw in iterator:
            values = list(raw)
            if not any(_cell_to_str(cell) for cell in values):
                continue
            padded = values + [None] * max(0, len(headers) - len(values))
            rows.append(
                {headers[i]: _cell_to_str(padded[i]) for i in range(len(headers))}
            )
        workbook.close()
        return headers, rows
    finally:
        tmp_path.unlink(missing_ok=True)


def parse_spreadsheet(content: bytes, filename: str) -> ParsedSpreadsheet:
    if not content:
        raise ValueError("Uploaded file is empty")

    name = (filename or "upload.csv").strip() or "upload.csv"
    lower = name.lower()
    file_hash = hashlib.sha256(content).hexdigest()

    is_csv = lower.endswith(".csv") or lower.endswith(".tsv")
    looks_ooxml = content[:2] == b"PK"
    is_excel = (
        lower.endswith(".xlsx")
        or lower.endswith(".xlsm")
        or lower.endswith(".xls")
        or looks_ooxml
    )

    if is_csv and not looks_ooxml:
        headers, rows = _parse_csv(content)
    elif is_excel:
        try:
            headers, rows = _parse_excel(content)
        except Exception as exc:
            # Misnamed CSV with an Excel extension.
            if is_csv or b"," in content[:2048]:
                headers, rows = _parse_csv(content)
            else:
                raise ValueError(f"Could not read spreadsheet: {exc}") from exc
    else:
        raise ValueError(
            "Unsupported file type. Upload a .csv, .xlsx, or .xls spreadsheet."
        )

    if not headers:
        raise ValueError("Spreadsheet has no header row")

    stats = _estimate_stats(headers, rows)
    samples = _build_samples(headers, rows)
    return ParsedSpreadsheet(
        source_file_name=name,
        source_file_hash=file_hash,
        headers=headers,
        rows=rows,
        samples=samples,
        stats=stats,
    )
